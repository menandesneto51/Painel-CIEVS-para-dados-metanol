from __future__ import annotations

import json
import os
import re
import sys
import tempfile
import unicodedata
from datetime import date, datetime
from pathlib import Path

import gdown
from openpyxl import load_workbook

DRIVE_FILE_ID = os.getenv("DRIVE_FILE_ID", "13sLXlRgOJFPWWxuRY6bbz0m8OJVG4-Gy")
SHEET_NAME = os.getenv("SHEET_NAME", "Cadastro de Casos")
INDEX_PATH = Path(os.getenv("INDEX_PATH", "index.html"))
MIN_CASES = int(os.getenv("MIN_CASES", "1"))
MAX_DROP_FRACTION = float(os.getenv("MAX_DROP_FRACTION", "0.25"))


def norm(v):
    if v is None:
        return ""
    s = str(v).strip()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", s).lower()


def fmt_date(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    return str(v).strip()


def valid_date_text(v):
    if not v:
        return False
    try:
        datetime.strptime(v, "%d/%m/%Y")
        return True
    except ValueError:
        return False


def ensure_dimensions(ws):
    if ws.max_row is None or ws.max_column is None:
        ws.calculate_dimension(force=True)


def find_header_row(ws):
    ensure_dimensions(ws)
    max_row = ws.max_row or 20
    max_col = ws.max_column or 30
    for r in range(1, min(max_row, 20) + 1):
        vals = [norm(ws.cell(r, c).value) for c in range(1, min(max_col, 30) + 1)]
        if "id do caso" in vals and "status do caso" in vals:
            return r
    raise RuntimeError("Cabeçalho da aba 'Cadastro de Casos' não encontrado.")


def build_header_map(ws, header_row):
    ensure_dimensions(ws)
    max_col = ws.max_column or 30
    return {norm(ws.cell(header_row, c).value): c for c in range(1, max_col + 1) if ws.cell(header_row, c).value}


def col(h, *names):
    for name in names:
        key = norm(name)
        if key in h:
            return h[key]
    return None


def get(ws, r, c):
    return ws.cell(r, c).value if c else None


def download_sheet(target: Path):
    result = gdown.download(id=DRIVE_FILE_ID, output=str(target), quiet=False)
    if not result or not target.exists() or target.stat().st_size < 1000:
        raise RuntimeError("Não foi possível baixar a planilha do Google Drive. Verifique se o arquivo está acessível por link.")


def read_cases(xlsx_path: Path):
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise RuntimeError(f"Aba obrigatória '{SHEET_NAME}' não encontrada. Abas disponíveis: {wb.sheetnames}")
    ws = wb[SHEET_NAME]
    ensure_dimensions(ws)
    hr = find_header_row(ws)
    h = build_header_map(ws, hr)

    c_id = col(h, "ID do caso")
    c_data = col(h, "Data da notificação")
    c_status = col(h, "Status do caso")
    c_res = col(h, "Município de residência")
    c_int = col(h, "Município de internação")
    c_sexo = col(h, "Sexo")
    c_idade = col(h, "Idade (anos)", "Idade")
    c_grav = col(h, "Gravidade")
    c_desc = col(h, "Descrição clínica/laboratorial")
    c_evol = col(h, "Evolução")
    c_hosp = col(h, "Necessitou hospitalização?")
    c_perm = col(h, "Permanece hospitalizado?")
    c_ant = col(h, "Recebeu antídoto?")
    c_obito = col(h, "Óbito?")
    c_tipo = col(h, "Tipo de óbito")
    c_data_obito = col(h, "Data do óbito")
    c_obs = col(h, "Observações")
    c_pronto = col(h, "Pronto para publicar?")

    required = {"ID do caso": c_id, "Data da notificação": c_data, "Status do caso": c_status, "Município de residência": c_res}
    missing = [k for k, v in required.items() if not v]
    if missing:
        raise RuntimeError("Colunas obrigatórias ausentes: " + ", ".join(missing))

    cases, skipped, seen_ids = [], [], set()
    max_row = ws.max_row or 0
    for r in range(hr + 1, max_row + 1):
        raw_id = get(ws, r, c_id)
        if raw_id in (None, ""):
            continue
        try:
            case_id = int(raw_id)
        except Exception:
            raise RuntimeError(f"ID inválido na linha {r}: {raw_id!r}")
        if case_id in seen_ids:
            raise RuntimeError(f"CONTROLE DE QUALIDADE: ID duplicado na planilha: {case_id}")
        seen_ids.add(case_id)

        pronto = norm(get(ws, r, c_pronto)) if c_pronto else ""
        if pronto in {"não", "nao"}:
            skipped.append(case_id)
            continue

        data_notificacao = fmt_date(get(ws, r, c_data))
        status = str(get(ws, r, c_status) or "").strip()
        residencia = str(get(ws, r, c_res) or "").strip()
        if not data_notificacao or not valid_date_text(data_notificacao):
            raise RuntimeError(f"CONTROLE DE QUALIDADE: data de notificação inválida no ID {case_id}: {data_notificacao!r}")
        if not status:
            raise RuntimeError(f"CONTROLE DE QUALIDADE: status vazio no ID {case_id}")
        if not residencia:
            raise RuntimeError(f"CONTROLE DE QUALIDADE: município de residência vazio no ID {case_id}")

        obito = str(get(ws, r, c_obito) or "Não informado").strip()
        tipo_obito = str(get(ws, r, c_tipo) or "Não se aplica").strip()
        data_obito = fmt_date(get(ws, r, c_data_obito))
        if norm(obito) == "sim":
            if norm(tipo_obito) in {"", "nao se aplica"}:
                raise RuntimeError(f"CONTROLE DE QUALIDADE: óbito sem tipo de óbito definido no ID {case_id}")
            if data_obito and not valid_date_text(data_obito):
                raise RuntimeError(f"CONTROLE DE QUALIDADE: data do óbito inválida no ID {case_id}: {data_obito!r}")
        # Não bloqueamos óbito anterior à notificação: notificações retrospectivas/pós-óbito são admitidas.

        idade = get(ws, r, c_idade)
        if isinstance(idade, float) and idade.is_integer():
            idade = int(idade)
        sexo = str(get(ws, r, c_sexo) or "Não informado").strip()

        cases.append({
            "id": case_id, "d": data_notificacao, "s": status, "r": residencia,
            "i": str(get(ws, r, c_int) or "Não informado").strip(), "x": sexo,
            "a": idade if idade not in (None, "") else "Não informado",
            "g": str(get(ws, r, c_grav) or "Não informado").strip(),
            "c": str(get(ws, r, c_desc) or "").strip(),
            "e": str(get(ws, r, c_evol) or "Não informado").strip(),
            "h": str(get(ws, r, c_hosp) or "Não informado").strip(),
            "p": str(get(ws, r, c_perm) or "Não informado").strip(),
            "t": str(get(ws, r, c_ant) or "Não informado").strip(),
            "o": obito, "to": tipo_obito, "do": data_obito,
            "obs": str(get(ws, r, c_obs) or "").strip(),
        })

    cases.sort(key=lambda x: x["id"])
    if len(cases) < MIN_CASES:
        raise RuntimeError(f"CONTROLE DE QUALIDADE: apenas {len(cases)} registro(s) apto(s); mínimo de segurança = {MIN_CASES}.")
    return cases, skipped


def current_panel_count(html):
    for pattern in (r"const C=(\[.*?\]);const Z=", r"const CASOS=(\[.*?\]);\s*const COORDS="):
        m = re.search(pattern, html, flags=re.S)
        if m:
            try:
                return len(json.loads(m.group(1)))
            except Exception:
                return None
    return None


def quality_gate(cases, html):
    old_count = current_panel_count(html)
    new_count = len(cases)
    if old_count and old_count > 0 and new_count < old_count:
        drop = (old_count - new_count) / old_count
        if drop > MAX_DROP_FRACTION:
            raise RuntimeError(
                f"CONTROLE DE QUALIDADE: redução abrupta de {old_count} para {new_count} registros ({drop:.1%}). "
                f"Limite automático = {MAX_DROP_FRACTION:.0%}. Publicação bloqueada; revisar a planilha."
            )


def update_html(cases):
    html = INDEX_PATH.read_text(encoding="utf-8")
    quality_gate(cases, html)
    payload = json.dumps(cases, ensure_ascii=False, separators=(",", ":"))
    patterns = [
        (r"const C=\[.*?\];const Z=", f"const C={payload};const Z="),
        (r"const CASOS=\[.*?\];\s*const COORDS=", f"const CASOS={payload};\nconst COORDS="),
    ]
    for pattern, repl in patterns:
        html, n = re.subn(pattern, repl, html, count=1, flags=re.S)
        if n == 1:
            break
    else:
        raise RuntimeError("Bloco de dados do painel não foi localizado no index.html.")

    today = datetime.now().astimezone().strftime("%d/%m/%Y")
    html = re.sub(r"Dados carregados da planilha atualizada em \d{2}/\d{2}/\d{4}\.", f"Dados carregados da planilha atualizada em {today}.", html, count=1)
    INDEX_PATH.write_text(html, encoding="utf-8")


def main():
    tmp = Path(tempfile.gettempdir()) / "metanol_drive.xlsx"
    if tmp.exists():
        tmp.unlink()
    download_sheet(tmp)
    cases, skipped = read_cases(tmp)
    update_html(cases)
    print(f"Controle de qualidade: APROVADO. Registros publicados: {len(cases)}")
    print("IDs publicados:", ", ".join(str(c["id"]) for c in cases))
    if skipped:
        print("IDs explicitamente não prontos e não publicados:", ", ".join(map(str, skipped)))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"ERRO: {exc}", file=sys.stderr)
        raise
